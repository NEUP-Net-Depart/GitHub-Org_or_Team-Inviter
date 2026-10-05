# -*- coding: utf-8 -*-
"""表格读取 (csv / tsv / xlsx) —— 零第三方依赖。

`.xlsx` 本质是个装了 XML 的 zip, 所以用标准库 `zipfile` + `xml.etree` 直接读;
装了 openpyxl 就优先用它(日期等格式更完整), 没装也不影响。
"""
from __future__ import annotations

import csv
import io
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

from .console import die, log
from .constants import TABLE_SUFFIXES


def _local(tag: str) -> str:
    """{namespace}row -> row"""
    return tag.rsplit("}", 1)[-1]


def _col_index(ref: str) -> int:
    """A1 -> 0, C2 -> 2"""
    letters = "".join(ch for ch in ref if ch.isalpha()).upper()
    index = 0
    for ch in letters:
        index = index * 26 + (ord(ch) - 64)
    return max(0, index - 1)


def _decode_bytes(raw: bytes) -> str:
    """按中文环境的现实情况依次尝试编码。"""
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "cp936", "latin-1"):
        try:
            return raw.decode(encoding)
        except (UnicodeDecodeError, LookupError):
            continue
    return raw.decode("utf-8", "replace")


def _guess_delimiter(text: str) -> str:
    """csv 可能是逗号/分号/制表符, 按第一行里出现次数最多的算。"""
    first_line = next((line for line in text.splitlines() if line.strip()), "")
    counts = {sep: first_line.count(sep) for sep in (",", ";", "\t")}
    best = max(counts, key=lambda sep: counts[sep])
    return best if counts[best] > 0 else ","


def _read_delimited(text: str, delimiter: str | None = None) -> list[list[str]]:
    sep = delimiter or _guess_delimiter(text)
    rows = [list(row) for row in csv.reader(io.StringIO(text), delimiter=sep)]
    cleaned = [[("" if cell is None else str(cell)).strip() for cell in row] for row in rows]
    # 完全空白的行(含末尾空行)直接丢掉: 它们不是数据, 报"排除了几行"只会让人困惑
    return [row for row in cleaned if any(cell for cell in row)]


def _xlsx_shared_strings(zf: zipfile.ZipFile) -> list[str]:
    try:
        root = ET.fromstring(zf.read("xl/sharedStrings.xml"))
    except (KeyError, ET.ParseError):
        return []
    out: list[str] = []
    for node in root:
        if _local(node.tag) != "si":
            continue
        out.append("".join(t.text or "" for t in node.iter() if _local(t.tag) == "t"))
    return out


def _xlsx_sheets(zf: zipfile.ZipFile) -> list[tuple[str, str]]:
    """返回 [(sheet 名, zip 内路径)]。"""
    try:
        workbook = ET.fromstring(zf.read("xl/workbook.xml"))
        rels = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
    except (KeyError, ET.ParseError):
        return [("Sheet1", "xl/worksheets/sheet1.xml")]

    rel_map: dict[str, str] = {}
    for rel in rels:
        rid = rel.get("Id") or ""
        target = (rel.get("Target") or "").lstrip("/")
        if not target:
            continue
        rel_map[rid] = target if target.startswith("xl/") else f"xl/{target}"

    out: list[tuple[str, str]] = []
    for node in workbook.iter():
        if _local(node.tag) != "sheet":
            continue
        rid = next((v for k, v in node.attrib.items() if _local(k) == "id"), "")
        out.append((node.get("name") or "", rel_map.get(rid or "", "xl/worksheets/sheet1.xml")))
    return out or [("Sheet1", "xl/worksheets/sheet1.xml")]


def _read_xlsx_minimal(path: Path, sheet_name: str = "") -> list[list[str]]:
    """不依赖 openpyxl 的最小 xlsx 读取器 —— xlsx 本质就是装了 XML 的 zip。"""
    try:
        with zipfile.ZipFile(path) as zf:
            shared = _xlsx_shared_strings(zf)
            sheets = _xlsx_sheets(zf)
            if sheet_name:
                target = next((t for name, t in sheets if name == sheet_name), None)
                if target is None:
                    die(f"工作表不存在: {sheet_name}\n"
                        f"现有工作表: {' | '.join(name for name, _ in sheets)}")
            else:
                target = sheets[0][1]
            try:
                root = ET.fromstring(zf.read(target))
            except KeyError:
                die(f"xlsx 内部缺少工作表数据: {target} (文件可能损坏)")
    except zipfile.BadZipFile:
        die(f"读不了这个 xlsx: {path.name} (不是有效的 xlsx, 或文件已损坏)")

    rows: list[list[str]] = []
    for row_node in root.iter():
        if _local(row_node.tag) != "row":
            continue
        values: dict[int, str] = {}
        highest = -1
        fallback = 0
        for cell in row_node:
            if _local(cell.tag) != "c":
                continue
            ref = cell.get("r") or ""
            ctype = cell.get("t") or "n"
            col = _col_index(ref) if ref else fallback
            fallback = col + 1
            text = ""
            if ctype == "inlineStr":
                text = "".join(n.text or "" for n in cell.iter() if _local(n.tag) == "t")
            else:
                vnode = next((c for c in cell if _local(c.tag) == "v"), None)
                if vnode is not None and vnode.text is not None:
                    if ctype == "s":
                        try:
                            text = shared[int(vnode.text)]
                        except (ValueError, IndexError):
                            text = vnode.text
                    elif ctype == "b":
                        text = "TRUE" if vnode.text == "1" else "FALSE"
                    else:
                        text = vnode.text
            values[col] = text.strip()
            highest = max(highest, col)
        rows.append([values.get(i, "") for i in range(highest + 1)] if highest >= 0 else [])
    # 和 csv 一样: 完全空白的行不是数据
    return [row for row in rows if any(cell for cell in row)]


def _read_xlsx(path: Path, sheet_name: str = "") -> list[list[str]]:
    """有 openpyxl 就用它(对日期等格式更稳), 没有就用自带的最小解析器。"""
    try:
        from openpyxl import load_workbook  # type: ignore
    except ImportError:
        return _read_xlsx_minimal(path, sheet_name)
    try:
        book = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - 解析失败就回退到最小解析器
        log(f"  ! openpyxl 读不了这个文件({type(exc).__name__}), 改用内置解析器")
        return _read_xlsx_minimal(path, sheet_name)
    try:
        if sheet_name:
            if sheet_name not in book.sheetnames:
                die(f"工作表不存在: {sheet_name}\n"
                    f"现有工作表: {' | '.join(book.sheetnames)}")
            sheet = book[sheet_name]
        else:
            sheet = book.active
        return [["" if cell is None else str(cell).strip() for cell in row]
                for row in sheet.iter_rows(values_only=True)]
    finally:
        book.close()


def read_table(path: str, sheet_name: str = "") -> list[list[str]]:
    """读 csv / xlsx。返回逐行的字符串矩阵。

    只支持表格文件: 一个纯文本清单每行只有一格, 装不下「邮箱 + 用户名」,
    与其为它单独维护一套解析语义, 不如要求用户给表格。
    """
    file_path = Path(path)
    if not file_path.is_file():
        die(f"找不到文件: {file_path}")
    suffix = file_path.suffix.lower()

    if suffix in (".xlsx", ".xlsm"):
        return _read_xlsx(file_path, sheet_name)
    if suffix in (".csv", ".tsv"):
        # .tsv 不算正式支持, 但它是分隔符明确的表格, 顺手认了
        text = _decode_bytes(file_path.read_bytes())
        return _read_delimited(text, "\t" if suffix == ".tsv" else None)
    if suffix == ".txt":
        die(f".txt 名单不再支持了: {file_path.name}\n"
            f"纯文本每行只有一格, 装不下「邮箱 + 用户名」两项, 也就分不出该用哪个邀请。\n"
            f"请改用表格(.csv 或 .xlsx), 列名写 邮箱 / GitHub用户名 即可。\n"
            f"示例见 docs/examples/")
    die(f"不支持的文件格式: {suffix or '(无扩展名)'}\n"
        f"支持: {', '.join(TABLE_SUFFIXES)}")
