#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
按数据表批量邀请成员加入 GitHub 组织。

依据 GitHub 官方接口:
    POST /orgs/{org}/invitations
    body: { "email": ... } 或 { "invitee_id": ... }, 可选 "role"、"team_ids"
    权限: 调用者必须是组织 owner;
          细粒度 PAT 需要 "Members" organization permissions (write);
          经典 PAT 需要 admin:org。

默认是 dry-run(只解析表格 + 预检, 不发任何邀请), 加 --execute 才会真正发送。

用法示例见同目录 README.md。
"""

from __future__ import annotations

import argparse
import csv
import io
import json
import os
import re
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
import xml.etree.ElementTree as ET
import zipfile
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

DEFAULT_API = "https://api.github.com"
API_VERSION = "2022-11-28"
USER_AGENT = "github-org-inviter/1.0"

# ---------------------------------------------------------------------------
# 表头识别
# ---------------------------------------------------------------------------

EMAIL_HEADERS = {
    "email", "emailaddress", "mail", "mailaddress", "e-mail", "electronicmail",
    "邮箱", "电子邮箱", "邮箱地址", "邮件", "电子邮件", "邮件地址", "电邮",
}
USERNAME_HEADERS = {
    "github", "githubusername", "githubname", "githubaccount", "githubid",
    "username", "login", "user", "handle", "ghusername",
    "用户名", "github用户名", "github账号", "账号", "githubid", "GitHub", "昵称账号",
}
NAME_HEADERS = {
    "name", "fullname", "realname", "nickname", "displayname",
    "姓名", "名字", "真实姓名", "昵称", "备注姓名",
}

EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s.]+(\.[^@\s.]+)+$")
LOGIN_RE = re.compile(r"^[A-Za-z0-9](?:[A-Za-z0-9]|-(?=[A-Za-z0-9])){0,38}$")


def norm_key(value: str) -> str:
    """把表头归一化, 便于匹配: 去空白/下划线/连字符/括号, 转小写。"""
    text = (value or "").strip().lower()
    text = re.sub(r"[\s_\-/\\()（）\[\]:：*]+", "", text)
    return text


def pick_column(header: list[str], wanted: set[str], keywords: Iterable[str]) -> int | None:
    keys = [norm_key(h) for h in header]
    wanted_norm = {norm_key(w) for w in wanted}
    for idx, key in enumerate(keys):
        if key and key in wanted_norm:
            return idx
    for idx, key in enumerate(keys):
        if key and any(k in key for k in keywords):
            return idx
    return None


def detect_columns(header: list[str]) -> dict[str, int | None]:
    return {
        "email": pick_column(header, EMAIL_HEADERS, ("邮箱", "email", "mail", "邮件")),
        "username": pick_column(header, USERNAME_HEADERS, ("github", "username", "login", "账号", "用户名")),
        "name": pick_column(header, NAME_HEADERS, ("姓名", "name", "昵称", "nickname")),
    }


# ---------------------------------------------------------------------------
# 表格读取
# ---------------------------------------------------------------------------

def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _col_index(ref: str) -> int:
    """A1 -> 0, C2 -> 2"""
    letters = "".join(ch for ch in ref if ch.isalpha()).upper()
    index = 0
    for ch in letters:
        index = index * 26 + (ord(ch) - 64)
    return max(0, index - 1)


def _xlsx_shared_strings(zf: zipfile.ZipFile) -> list[str]:
    try:
        root = ET.fromstring(zf.read("xl/sharedStrings.xml"))
    except KeyError:
        return []
    out: list[str] = []
    for si in root:
        if _local(si.tag) != "si":
            continue
        out.append("".join(node.text or "" for node in si.iter() if _local(node.tag) == "t"))
    return out


def _xlsx_sheet_path(zf: zipfile.ZipFile, sheet_name: str) -> str:
    try:
        wb = ET.fromstring(zf.read("xl/workbook.xml"))
        rels = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
    except KeyError:
        return "xl/worksheets/sheet1.xml"
    rel_map = {rel.get("Id"): rel.get("Target") for rel in rels}
    sheets: list[tuple[str, str | None]] = []
    for node in wb.iter():
        if _local(node.tag) != "sheet":
            continue
        rid = next((v for k, v in node.attrib.items() if _local(k) == "id"), None)
        sheets.append((node.get("name") or "", rid))
    if not sheets:
        return "xl/worksheets/sheet1.xml"
    if sheet_name:
        match = next((rid for name, rid in sheets if name == sheet_name), None)
        if match is None and sheet_name not in [n for n, _ in sheets]:
            die(f"工作表不存在: {sheet_name}\n现有工作表: {' | '.join(n for n, _ in sheets)}")
        target = rel_map.get(match or "")
    else:
        target = rel_map.get(sheets[0][1] or "")
    target = (target or "worksheets/sheet1.xml").lstrip("/")
    return target if target.startswith("xl/") else f"xl/{target}"


def _read_xlsx_minimal(path: Path, sheet_name: str = "") -> list[list[str]]:
    """不依赖 openpyxl 的最小 xlsx 读取器: xlsx 本质是装了 XML 的 zip。"""
    with zipfile.ZipFile(path) as zf:
        shared = _xlsx_shared_strings(zf)
        target = _xlsx_sheet_path(zf, sheet_name)
        try:
            root = ET.fromstring(zf.read(target))
        except KeyError:
            die(f"xlsx 内部缺少工作表数据: {target} (文件可能损坏)")
    rows: list[list[str]] = []
    for row_node in root.iter():
        if _local(row_node.tag) != "row":
            continue
        values: dict[int, str] = {}
        highest = -1
        for cell in row_node:
            if _local(cell.tag) != "c":
                continue
            ref = cell.get("r") or ""
            ctype = cell.get("t") or "n"
            col = _col_index(ref) if ref else highest + 1
            text = ""
            if ctype == "inlineStr":
                text = "".join(n.text or "" for n in cell.iter() if _local(n.tag) == "t")
            else:
                vnode = next((c for c in cell if _local(c.tag) == "v"), None)
                if vnode is not None and vnode.text is not None:
                    if ctype == "s":
                        try:
                            text = shared[int(vnode.text)]
                        except (ValueError, IndexError):
                            text = vnode.text
                    elif ctype == "b":
                        text = "TRUE" if vnode.text == "1" else "FALSE"
                    else:
                        text = vnode.text
            values[col] = text
            highest = max(highest, col)
        rows.append([values.get(i, "") for i in range(highest + 1)] if highest >= 0 else [])
    return rows


def read_table(path: Path, sheet_name: str = "") -> list[list[str]]:
    suffix = path.suffix.lower()
    if suffix in (".csv", ".tsv", ".txt"):
        delimiter = "\t" if suffix == ".tsv" else ","
        raw = path.read_bytes()
        for encoding in ("utf-8-sig", "utf-8", "gb18030", "cp936", "latin-1"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:  # pragma: no cover - 兜底
            text = raw.decode("utf-8", "replace")
        # 自动识别分隔符(csv 可能是分号)
        if suffix == ".csv":
            sample = text[:4000]
            delimiter = max((",", ";", "\t"), key=sample.count)
        return [row for row in csv.reader(io.StringIO(text), delimiter=delimiter)]
    if suffix in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook  # type: ignore
        except ImportError:
            return _read_xlsx_minimal(path, sheet_name)
        book = load_workbook(path, read_only=True, data_only=True)
        if sheet_name:
            if sheet_name not in book.sheetnames:
                die(f"工作表不存在: {sheet_name}\n现有工作表: {' | '.join(book.sheetnames)}")
            sheet = book[sheet_name]
        else:
            sheet = book.active
        rows: list[list[str]] = []
        for row in sheet.iter_rows(values_only=True):
            rows.append(["" if cell is None else str(cell) for cell in row])
        book.close()
        return rows
    die(f"不支持的表格格式: {suffix} (支持 .csv / .tsv / .xlsx)")


# ---------------------------------------------------------------------------
# 记录
# ---------------------------------------------------------------------------

@dataclass
class Record:
    row: int
    email: str = ""
    login: str = ""
    name: str = ""
    note: str = ""

    def key(self) -> str:
        return (self.email or self.login).lower()


@dataclass
class Outcome:
    record: Record
    status: str
    http_status: int = 0
    detail: str = ""
    invitee_id: int | None = None


def clean_email(value: str) -> tuple[str, str]:
    """返回 (清洗后的邮箱, 备注)。单元格里有多个邮箱时只取第一个。"""
    text = (value or "").strip()
    if not text:
        return "", ""
    text = text.replace("\u3000", " ")
    text = re.sub(r"^mailto:", "", text, flags=re.I)
    parts = [p.strip().strip("<>").strip() for p in re.split(r"[,;、/|\s]+", text) if p.strip()]
    valid = [p for p in parts if EMAIL_RE.match(p)]
    if not valid:
        return "", f"无有效邮箱: {text[:60]}"
    note = ""
    if len(valid) > 1:
        note = f"单元格含 {len(valid)} 个邮箱, 只取第一个"
    return valid[0].lower(), note


def clean_login(value: str) -> str:
    text = (value or "").strip()
    if not text:
        return ""
    # 兼容 https://github.com/x、github.com/x、www.github.com/x 这几种写法
    text = re.sub(r"^(?:https?://)?(?:www\.)?github\.com/", "", text, flags=re.I)
    text = text.strip().strip("@/").strip()
    text = text.split("/")[0].split("?")[0]
    return text


def classify_cell(value: str) -> tuple[str, str]:
    """把一个单元格判成 (邮箱, 用户名), 两者最多填一个。

    这样即便两列写反了(邮箱列里写用户名、用户名列里写邮箱), 或者干脆
    只有一列把两种混着写, 也能各归各位 —— 不靠列名, 靠内容。
    """
    text = (value or "").strip()
    if not text:
        return "", ""
    email, _ = clean_email(text)
    if email:
        return email, ""
    login = clean_login(text)
    if not login:
        return "", ""
    if EMAIL_RE.match(login):
        return login.lower(), ""
    return "", login


ROLE_ALIASES = {
    "username": "username", "user": "username", "login": "username", "用户名": "username",
    "email": "email", "mail": "email", "邮箱": "email",
    "name": "name", "姓名": "name", "名字": "name",
}


def parse_cols_spec(spec: str) -> dict[str, int]:
    """解析 --cols: 按**位置**规定每列是什么。

    例如 "username,email" 表示第 1 列是用户名、第 2 列是邮箱。
    某一列不想用就写 "-" 占位, 例如 "-,username,email"(跳过第 1 列)。
    """
    out: dict[str, int] = {}
    for idx, token in enumerate(str(spec or "").split(",")):
        key = token.strip().lower()
        if not key or key in ("-", "_", "skip", "忽略", "无"):
            continue
        role = ROLE_ALIASES.get(key)
        if role:
            out[role] = idx
        else:
            die(f'--cols 第 {idx + 1} 项 "{token.strip()}" 看不懂。\n'
                f'每项只能填 username / email / name, 或 "-" 占位。\n'
                f'例如: --cols "username,email"')
    return out


def build_records(rows: list[list[str]], cols: dict[str, int | None], sheet_name: str) -> list[Record]:
    if not rows:
        die(f"{sheet_name}: 表格是空的。")
    header = [str(c) for c in rows[0]]
    if cols["email"] is None and cols["username"] is None:
        die("没识别出邮箱列或 GitHub 用户名列。表头是:\n    "
            + " | ".join(header)
            + "\n请用 --cols \"username,email\" 按位置指定, 或 --email-col / --username-col 指定列名。")

    records: list[Record] = []
    for offset, row in enumerate(rows[1:], start=2):
        def cell(idx: int | None) -> str:
            if idx is None or idx >= len(row):
                return ""
            return str(row[idx]).strip()

        raw_email = cell(cols["email"])
        raw_user = cell(cols["username"])
        name = cell(cols["name"])

        # 按「内容」判断而不是按列名 —— 两列写反、或只有一列混装, 都能各归各位。
        # 优先信各自的本职列, 另一列只作兜底: 否则邮箱列里的垃圾值(比如
        # "not-an-email", 恰好是合法用户名格式)会盖掉用户名列里真正填的值。
        email_a, login_a = classify_cell(raw_email)
        email_b, login_b = classify_cell(raw_user)
        email = email_a or email_b
        login = login_b or login_a

        note = ""
        if raw_user and not login_b and not email_b:
            note = f"用户名列内容无法识别: {raw_user[:40]}"
        elif raw_email and not email_a and not login_a:
            note = f"邮箱列内容无法识别: {raw_email[:40]}"
        elif login and not LOGIN_RE.match(login):
            note = f"用户名格式可疑: {login}"

        if not email and not login and not name:
            continue  # 跳过空行
        records.append(Record(row=offset, email=email, login=login, name=name, note=note))
    return records


# ---------------------------------------------------------------------------
# GitHub API
# ---------------------------------------------------------------------------

class ApiError(Exception):
    def __init__(self, status: int, message: str, payload: Any = None):
        super().__init__(message)
        self.status = status
        self.message = message
        self.payload = payload


class GitHub:
    def __init__(self, token: str, api_url: str = DEFAULT_API, timeout: int = 30, verbose: bool = False):
        self.token = token
        self.api_url = api_url.rstrip("/")
        self.timeout = timeout
        self.verbose = verbose

    def _once(self, method: str, path: str, body: dict | None = None) -> tuple[int, Any, dict]:
        url = path if path.startswith("http") else f"{self.api_url}{path}"
        data = json.dumps(body).encode("utf-8") if body is not None else None
        req = urllib.request.Request(url, data=data, method=method)
        req.add_header("Accept", "application/vnd.github+json")
        req.add_header("Authorization", f"Bearer {self.token}")
        req.add_header("X-GitHub-Api-Version", API_VERSION)
        req.add_header("User-Agent", USER_AGENT)
        if data is not None:
            req.add_header("Content-Type", "application/json")
        try:
            with urllib.request.urlopen(req, timeout=self.timeout) as resp:
                raw = resp.read().decode("utf-8", "replace")
                payload = json.loads(raw) if raw.strip() else {}
                return resp.status, payload, dict(resp.headers)
        except urllib.error.HTTPError as exc:
            raw = exc.read().decode("utf-8", "replace")
            try:
                payload = json.loads(raw) if raw.strip() else {}
            except json.JSONDecodeError:
                payload = {"message": raw[:400]}
            return exc.code, payload, dict(exc.headers or {})
        except urllib.error.URLError as exc:
            raise ApiError(0, f"网络错误: {exc.reason}") from exc
        except TimeoutError as exc:
            raise ApiError(0, "请求超时") from exc

    def request(self, method: str, path: str, body: dict | None = None,
                max_retries: int = 4) -> tuple[int, Any]:
        """带限流重试的请求。返回 (status, payload); status>=400 时也正常返回。"""
        attempt = 0
        while True:
            attempt += 1
            try:
                status, payload, headers = self._once(method, path, body)
            except ApiError:
                if attempt > max_retries:
                    raise
                wait = min(60, 2 ** attempt)
                log(f"  ! 网络异常, {wait}s 后重试 ({attempt}/{max_retries})")
                time.sleep(wait)
                continue

            if status < 400 or status not in (403, 429, 500, 502, 503, 504):
                return status, payload

            message = str(payload.get("message", "")) if isinstance(payload, dict) else ""
            lowered = message.lower()
            wait: float | None = None

            retry_after = headers.get("Retry-After")
            if retry_after and retry_after.isdigit():
                wait = float(retry_after)

            if status == 403 and "secondary rate limit" in lowered:
                wait = wait or min(120.0, 60.0 * attempt)
                log(f"  ! 触发二级限流, 等待 {wait:.0f}s")
            elif status == 403 and headers.get("X-RateLimit-Remaining") == "0":
                reset = headers.get("X-RateLimit-Reset")
                wait = max(wait or 0.0, (float(reset) - time.time() + 2) if reset else 60.0)
                wait = max(wait, 1.0)
                log(f"  ! 主限流已用尽, 等待 {wait:.0f}s")
            elif status in (500, 502, 503, 504):
                wait = wait or min(30.0, 3.0 * attempt)
            elif status == 429:
                wait = wait or 30.0

            if wait is None or attempt > max_retries:
                return status, payload

            time.sleep(min(wait, 900.0))

    def get_paged(self, path: str, limit: int = 0) -> list[Any]:
        items: list[Any] = []
        page = 1
        sep = "&" if "?" in path else "?"
        while True:
            status, payload = self.request("GET", f"{path}{sep}per_page=100&page={page}")
            if status >= 400:
                raise ApiError(status, _describe(status, payload), payload)
            if not isinstance(payload, list) or not payload:
                break
            items.extend(payload)
            if len(payload) < 100 or (limit and len(items) >= limit):
                break
            page += 1
        return items[:limit] if limit else items


def _describe(status: int, payload: Any) -> str:
    if isinstance(payload, dict):
        message = str(payload.get("message") or "").strip()
        errors = payload.get("errors")
        if isinstance(errors, list) and errors:
            parts = []
            for err in errors:
                if isinstance(err, dict):
                    parts.append(str(err.get("message") or err.get("field") or err))
                else:
                    parts.append(str(err))
            message = f"{message} ({'; '.join(parts)})" if message else "; ".join(parts)
        if message:
            return message
    return f"HTTP {status}"


# ---------------------------------------------------------------------------
# 输出
# ---------------------------------------------------------------------------

def log(message: str = "") -> None:
    print(message, flush=True)


def die(message: str, code: int = 2) -> None:
    print(f"错误: {message}", file=sys.stderr, flush=True)
    raise SystemExit(code)


def now_iso() -> str:
    return datetime.now(timezone.utc).astimezone().strftime("%Y-%m-%d %H:%M:%S")


# ---------------------------------------------------------------------------
# 主流程
# ---------------------------------------------------------------------------

def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="按数据表批量邀请成员加入 GitHub 组织(默认 dry-run)",
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    parser.add_argument("--org", required=True, help="组织名(slug), 例如 my-org")
    parser.add_argument("--input", default="", help="数据表路径(.csv/.tsv/.xlsx); --report 时可以省略")
    parser.add_argument("--execute", action="store_true",
                        help="真正发送邀请。不加这个参数就是 dry-run, 不会发出任何邀请。")
    parser.add_argument("--role", default="direct_member",
                        choices=["direct_member", "admin", "billing_manager"],
                        help="邀请角色, 默认 direct_member")
    parser.add_argument("--team-ids", default="",
                        help="逗号分隔的 team id, 例如 12,26")
    parser.add_argument("--email-col", default="", help="手动指定邮箱列名")
    parser.add_argument("--username-col", default="", help="手动指定 GitHub 用户名列名")
    parser.add_argument("--cols", default="",
                        help='按位置规定每列是什么, 例如 "username,email" '
                             '= 第1列用户名、第2列邮箱(不想用的列写 "-")')
    parser.add_argument("--sheet", default="", help="xlsx 工作表名, 默认第一个")
    parser.add_argument("--delay", type=float, default=2.0,
                        help="每次发邀请之间的间隔秒数, 默认 2.0")
    parser.add_argument("--limit", type=int, default=0,
                        help="本次最多处理多少条(0=不限制), 便于小批量试跑")
    parser.add_argument("--prefer-username", action="store_true", default=True,
                        help="有用户名时优先用 invitee_id 邀请(默认开启, 更可靠)")
    parser.add_argument("--prefer-email", dest="prefer_username", action="store_false",
                        help="强制只按邮箱邀请")
    parser.add_argument("--no-preflight", action="store_true",
                        help="跳过预检(不查询已有成员/待处理邀请)")
    parser.add_argument("--no-resume", action="store_true",
                        help="忽略上次结果, 重新处理全部记录(可能重复发送)")
    parser.add_argument("--out-dir", default="", help="结果输出目录, 默认放脚本同级 output/")
    parser.add_argument("--api-url", default=DEFAULT_API, help="GitHub API 地址")
    parser.add_argument("--token", default="",
                        help="PAT。优先级: --token > 环境变量 GITHUB_TOKEN > 脚本目录的 token.txt")
    parser.add_argument("--report", action="store_true",
                        help="对账模式: 拉取组织的待处理/失败邀请, 和表格逐条比对(带上 --input 效果最好)")
    parser.add_argument("--resolve-member-emails", action="store_true",
                        help="预检时尝试读取每位成员的公开邮箱, 用来提前认出已经是成员的人(较慢)")
    parser.add_argument("--resolve-emails", action="store_true",
                        help="先用搜索接口把邮箱反查成 GitHub 用户名(命中就能精确判断是否已在组织内, 并改走用户名邀请)")
    parser.add_argument("--verbose", action="store_true", help="输出更多细节")
    return parser.parse_args(argv)


TOKEN_FILE_NAMES = ("token.txt", ".env")
TOKEN_KEYS = ("GITHUB_TOKEN", "GH_TOKEN", "TOKEN")


def _read_token_file(path: Path) -> str:
    """从 token.txt / .env 里读 token。支持 KEY=VALUE, 也支持整份文件只写一个 token。"""
    try:
        text = path.read_text(encoding="utf-8-sig", errors="replace")
    except OSError:
        return ""
    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#"):
            continue
        if "=" in line:
            key, _, value = line.partition("=")
            if key.strip().upper() not in TOKEN_KEYS:
                continue
        else:
            value = line
        value = value.strip().strip('"').strip("'").strip()
        if value:
            return value
    return ""


def find_token_file() -> tuple[str, Path | None]:
    """先找当前目录, 再找脚本所在目录, 避免你每次都得 cd 对位置。"""
    here = Path(__file__).resolve().parent
    for directory in (Path.cwd(), here):
        for name in TOKEN_FILE_NAMES:
            candidate = directory / name
            if candidate.is_file():
                token = _read_token_file(candidate)
                if token:
                    return token, candidate
    return "", None


def resolve_token(args: argparse.Namespace, required: bool) -> str:
    token = (args.token or "").strip()
    source = "--token 参数"
    if not token:
        token = (os.environ.get("GITHUB_TOKEN") or os.environ.get("GH_TOKEN") or "").strip()
        source = "环境变量"
    if not token:
        token, path = find_token_file()
        if path is not None:
            source = f"配置文件 {path.name}"
    if not token and required:
        die("没找到 token。三选一:\n"
            "  1) 在脚本目录建一个 token.txt, 里面写一行 GITHUB_TOKEN=ghp_xxx  (推荐, 不用每次设)\n"
            "  2) 设环境变量:  PowerShell  $env:GITHUB_TOKEN = \"ghp_xxx\"\n"
            "                 Git Bash    export GITHUB_TOKEN=\"ghp_xxx\"\n"
            "  3) 临时用 --token 传(会留在命令历史里, 不推荐)\n"
            "提示: 只是想先看看解析结果的话, 加 --no-preflight 就能免 token 跑 dry-run。")
    if token:
        log(f"token 来源: {source}")
    return token


def load_done_keys(out_dir: Path) -> set[str]:
    """扫描历史结果, 收集已经处理成功过的邮箱/用户名, 供续跑跳过。"""
    done: set[str] = set()
    for path in sorted(out_dir.glob("results-*.jsonl")):
        try:
            # utf-8-sig: 兼容带 BOM 的文件(比如用记事本编辑过), 否则第一行会解析失败被丢掉
            lines = path.read_text(encoding="utf-8-sig", errors="replace").splitlines()
        except OSError:
            continue
        for line in lines:
            if not line.strip():
                continue
            try:
                item = json.loads(line)
            except json.JSONDecodeError:
                continue
            if item.get("status") in ("invited", "already_member", "already_invited"):
                for key in ("email", "login"):
                    value = str(item.get(key) or "").strip().lower()
                    if value:
                        done.add(value)
    return done


def load_records(args: argparse.Namespace) -> tuple[list[Record], list[str], dict[str, int | None], Path, int]:
    table_path = Path(args.input)
    if not table_path.is_file():
        die(f"找不到数据表: {table_path}")
    rows = read_table(table_path, args.sheet)
    header = [str(c) for c in rows[0]] if rows else []
    cols = detect_columns(header)
    if args.email_col:
        cols["email"] = 0 if args.email_col.isdigit() else _index_of(header, args.email_col)
    if args.username_col:
        cols["username"] = 0 if args.username_col.isdigit() else _index_of(header, args.username_col)
    if getattr(args, "cols", ""):
        spec_cols = parse_cols_spec(args.cols)
        width = max((len(r) for r in rows), default=0)
        for role, idx in sorted(spec_cols.items(), key=lambda kv: kv[1]):
            if idx >= width:
                die(f'--cols 里 {role} 指的是第 {idx + 1} 列, 但这张表只有 {width} 列。'
                    f'\n(提醒: 位置是从 1 开始数的, "username,email" = 第1列用户名、第2列邮箱)')
        cols.update(spec_cols)
    if cols["email"] is None and cols["username"] is None:
        # 表头认不出来时, 按约定位置: 第 1 列用户名, 第 2 列邮箱
        if len(header) >= 1:
            cols["username"] = 0
        if len(header) >= 2:
            cols["email"] = 1
        log('  表头没认出来, 按约定位置处理: 第 1 列=用户名, 第 2 列=邮箱'
            ' (要改就加 --cols "username,email")')
    records = build_records(rows, cols, table_path.name)
    return records, header, cols, table_path, max(0, len(rows) - 1)


def resolve_emails_to_logins(gh: GitHub, records: list[Record]) -> dict[str, str]:
    """用搜索接口把邮箱反查成 GitHub 登录名。

    只对「把邮箱设为公开」的账号有效, 所以命中率有限; 但一旦命中就赚了:
      - 能拿登录名去比成员列表, 精确判断这人是不是已经在组织里
      - 还能改走 invitee_id 邀请, 绕开「邮箱必须是已验证邮箱」的限制
    搜索接口限速 30 次/分钟, 所以每条之间 sleep 2 秒。
    """
    resolved: dict[str, str] = {}
    for rec in records:
        query = urllib.parse.quote(f"{rec.email} in:email", safe="")
        code, payload = gh.request("GET", f"/search/users?q={query}")
        if code < 400 and isinstance(payload, dict):
            for item in (payload.get("items") or [])[:3]:
                login = str(item.get("login") or "")
                if not login:
                    continue
                # 二次确认: 拉一次资料, 公开邮箱必须完全一致, 避免模糊匹配误判
                check, profile = gh.request("GET", f"/users/{urllib.parse.quote(login)}")
                if check < 400 and isinstance(profile, dict) \
                        and str(profile.get("email") or "").strip().lower() == rec.email.lower():
                    resolved[rec.email.lower()] = login
                    break
        time.sleep(2)
    return resolved


def load_attempts(out_dir: Path) -> dict[str, dict]:
    """读历史结果文件, 弄清「哪些邮箱/用户名真的被发过、那次接口是怎么回的」。"""
    attempts: dict[str, dict] = {}
    for path in sorted(out_dir.glob("results-*.jsonl")):
        try:
            lines = path.read_text(encoding="utf-8-sig", errors="replace").splitlines()
        except OSError:
            continue
        for line in lines:
            if not line.strip():
                continue
            try:
                item = json.loads(line)
            except json.JSONDecodeError:
                continue
            for field in ("email", "login"):
                key = str(item.get(field) or "").strip().lower()
                if key:
                    attempts[key] = item  # 越晚的记录越新, 覆盖旧的
    return attempts


def apply_email_resolution(gh: GitHub, records: list[Record]) -> dict[str, str]:
    """把邮箱反查成登录名, 就地写回记录。发送路径和对账路径都要用。"""
    targets = [r for r in records if r.email and not r.login]
    if not targets:
        log("\n反查邮箱: 没有需要反查的行(都已经有用户名了)。")
        return {}
    log(f"\n反查邮箱: 尝试把 {len(targets)} 个邮箱换成 GitHub 登录名 ...")
    if len(targets) > 60:
        log(f"  注意: 搜索接口限速 30 次/分钟, {len(targets)} 条约需 {len(targets) * 2 // 60} 分钟。")
    found = resolve_emails_to_logins(gh, targets)
    for rec in records:
        login = found.get(rec.email.lower())
        if login and not rec.login:
            rec.login = login
            rec.note = (rec.note + "; " if rec.note else "") + "由邮箱反查得到登录名"
    log(f"  反查到 {len(found)} / {len(targets)} 个。")
    for email, login in found.items():
        log(f"    {email} -> {login}")
    if not found:
        log("  (多数人没在 GitHub 上公开邮箱, 反查不到是正常的)")
    return found


def reconcile(gh: GitHub, org: str, records: list[Record], out_dir: Path) -> None:
    """把「我们实际发过什么」和「组织当前的真实状态」对起来。

    关键: 绝不能把「还没发过」误判成「已经是成员」。
    所以先读 results-*.jsonl 确认每一行到底有没有被尝试过, 再看接口当时的回复。
    """
    log(f"== 对账: {org} ==")
    attempts = load_attempts(out_dir)
    if not attempts:
        log("  提示: 还没找到任何发送记录。对账要先跑过 --execute 才有意义。")

    try:
        pending = gh.get_paged(f"/orgs/{org}/invitations")
        failed = gh.get_paged(f"/orgs/{org}/failed_invitations")
        members = gh.get_paged(f"/orgs/{org}/members")
    except ApiError as exc:
        die(f"拉取组织数据失败: {exc.message}\n"
            "检查: 组织名是否正确、token 是否有效、是否有该组织的 owner 权限。")

    pending_emails = {str(i.get("email") or "").strip().lower() for i in pending if i.get("email")}
    pending_logins = {str(i.get("login") or "").strip().lower() for i in pending if i.get("login")}
    failed_map = {
        str(i.get("email") or "").strip().lower(): str(i.get("failed_reason") or "未给出原因")
        for i in failed if i.get("email")
    }
    member_logins = {str(m.get("login") or "").strip().lower() for m in members if m.get("login")}

    log(f"组织现有成员 {len(member_logins)} 人, 待处理邀请 {len(pending)} 个, "
        f"失败邀请 {len(failed)} 个。")
    if failed:
        log("  组织当前的失败邀请(可能是历史上别人发的):")
        for item in failed[:20]:
            who = item.get("email") or item.get("login") or "?"
            log(f"    {who}  —— {item.get('failed_reason') or '未给出原因'}")
    log("")

    rows_out: list[list[str]] = []
    counts: dict[str, int] = {}
    for rec in records:
        key_email = rec.email.lower()
        key_login = rec.login.lower()
        attempt = attempts.get(key_email) or attempts.get(key_login)

        if attempt is None:
            if key_login and key_login in member_logins:
                verdict, detail = "已经是组织成员", "在成员列表里(本次没给它发邀请)"
            else:
                verdict, detail = "未处理", "还没给它发过邀请(可能被 --limit 截断, 或还没跑)"
        else:
            status = str(attempt.get("status") or "")
            api_detail = str(attempt.get("detail") or "")
            # 用那次发送时记录下来的身份去比对, 比用表格里的更准
            a_email = str(attempt.get("email") or "").strip().lower()
            a_login = str(attempt.get("login") or "").strip().lower()
            if status == "already_member":
                verdict, detail = "已经是组织成员", api_detail or "接口告知此人已是成员"
            elif status == "invited":
                if (a_email and a_email in pending_emails) or (a_login and a_login in pending_logins):
                    verdict, detail = "邀请已发出", "等待对方接受"
                elif a_login and a_login in member_logins:
                    verdict, detail = "已经加入", "对方已接受邀请"
                else:
                    verdict, detail = "已发出, 但不在待处理", \
                        "很可能是对方已经接受加入了(去组织成员页确认); 也可能邀请没真正建成"
            elif status == "failed":
                verdict, detail = "邀请失败", api_detail
            elif a_email and a_email in failed_map:
                verdict, detail = "邀请失败", failed_map[a_email]
            elif a_login and a_login in member_logins:
                verdict, detail = "已经是组织成员", "在成员列表里"
            else:
                verdict, detail = status or "未知", api_detail

        counts[verdict] = counts.get(verdict, 0) + 1
        who = rec.login or rec.email or rec.name
        log(f"  {rec.row:>3} 行  {who:<28} {verdict}"
            + (f"  ({detail})" if detail else ""))
        rows_out.append([rec.row, rec.name, rec.email, rec.login, verdict, detail])

    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    out_path = out_dir / f"verify-{stamp}.csv"
    with out_path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow(["row", "name", "email", "username", "verdict", "detail"])
        writer.writerows(rows_out)

    log("\n== 汇总 ==")
    for verdict, count in sorted(counts.items(), key=lambda kv: -kv[1]):
        log(f"  {verdict}: {count}")
    log(f"\n对账结果: {out_path}")
    log("「未处理」= 还没发过, 不是已经是成员。要发剩下的就加 --execute 再跑一次。")


def do_report(gh: GitHub, org: str) -> None:
    log(f"== 组织 {org} 的邀请状态 ==")
    sections = [
        ("待处理邀请 (pending)", f"/orgs/{org}/invitations"),
        ("失败邀请 (failed)", f"/orgs/{org}/failed_invitations"),
    ]
    for title, path in sections:
        log(f"\n-- {title} --")
        try:
            items = gh.get_paged(path)
        except ApiError as exc:
            log(f"  查询失败: {exc.message}")
            continue
        if not items:
            log("  (空)")
            continue
        for item in items:
            login = item.get("login") or "(未知用户名)"
            email = item.get("email") or "-"
            extra = ""
            if item.get("failed_reason"):
                extra = f"  失败原因: {item['failed_reason']}"
            log(f"  {login:<24} {email}{extra}")


def main(argv: list[str] | None = None) -> int:
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")  # type: ignore[attr-defined]
    except Exception:
        pass

    args = parse_args(argv)
    if not args.report and not args.input:
        die("必须提供 --input(数据表路径)。")
    needs_token = args.execute or args.report or not args.no_preflight or args.resolve_emails
    token = resolve_token(args, required=needs_token)
    gh = GitHub(token, args.api_url, verbose=args.verbose) if token else None

    out_dir = Path(args.out_dir) if args.out_dir else Path(__file__).resolve().parent / "output"
    out_dir.mkdir(parents=True, exist_ok=True)

    if args.report:
        assert gh is not None
        if args.input:
            records, _header, _cols, _path, _total = load_records(args)
            if args.resolve_emails:
                apply_email_resolution(gh, records)
            reconcile(gh, args.org, records, out_dir)
        else:
            do_report(gh, args.org)
        return 0

    records, header, cols, table_path, total_rows = load_records(args)

    log(f"数据表: {table_path}")
    log(f"表头: {' | '.join(header)}")
    log(f"列映射: 邮箱=<{_name_of(header, cols['email'])}>  "
        f"用户名=<{_name_of(header, cols['username'])}>  "
        f"姓名=<{_name_of(header, cols['name'])}>")

    # ---- 去重 / 校验 ----
    valid: list[Record] = []
    excluded: list[Outcome] = []
    warnings: list[str] = []
    seen: dict[str, int] = {}
    for rec in records:
        if not rec.email and not rec.login:
            excluded.append(Outcome(rec, "skipped_invalid", detail=rec.note or "既没有邮箱也没有用户名"))
            continue
        if rec.note:
            warnings.append(f"第 {rec.row} 行: {rec.note}")
        if not rec.email:
            warnings.append(f"第 {rec.row} 行: 没有可用邮箱, 只能按用户名 {rec.login} 邀请")
        key = rec.key()
        if key in seen:
            excluded.append(Outcome(rec, "skipped_duplicate", detail=f"与第 {seen[key]} 行重复"))
            continue
        seen[key] = rec.row
        valid.append(rec)

    blank_rows = total_rows - len(records)
    log(f"\n解析到 {len(records)} 行数据"
        + (f"(另有 {blank_rows} 行空白已忽略)" if blank_rows else "")
        + f", 其中 {len(excluded)} 行被排除, 待处理 {len(valid)} 条。")
    for out in excluded:
        log(f"  - 排除第 {out.record.row} 行: {out.detail}")
    for item in warnings:
        log(f"  ! 注意 {item}")

    if not valid:
        die("没有可处理的记录。")

    # ---- 邮箱反查用户名(可选): 越早做越好, 反查到的登录名能参与后面的成员比对 ----
    if args.resolve_emails:
        assert gh is not None
        apply_email_resolution(gh, valid)

    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")

    # ---- 预检: 已有成员 / 待处理邀请 ----
    members: set[str] = set()
    member_emails: set[str] = set()
    pending_logins: set[str] = set()
    pending_emails: set[str] = set()
    pending_total = 0
    if not args.no_preflight:
        assert gh is not None
        log("\n预检: 读取组织已有成员与待处理邀请 ...")
        try:
            for item in gh.get_paged(f"/orgs/{args.org}/members"):
                if item.get("login"):
                    members.add(item["login"].lower())
            for item in gh.get_paged(f"/orgs/{args.org}/invitations"):
                pending_total += 1          # 按邮箱发的邀请 login 是 null, 不能只数 logins
                if item.get("login"):
                    pending_logins.add(item["login"].lower())
                if item.get("email"):
                    pending_emails.add(item["email"].lower())
        except ApiError as exc:
            log(f"  预检失败: {exc.message}")
            if exc.status in (401, 403, 404):
                die("预检无法完成, 说明 token 权限或组织名有问题。请先修好再跑。")
            log("  继续, 但可能产生重复邀请。")
        log(f"  已有成员 {len(members)} 人, 待处理邀请 {pending_total} 个。")

        # 成员列表不含邮箱, 只能靠大家公开资料里的 email 尽力而为
        if args.resolve_member_emails and members:
            log(f"  尝试解析 {len(members)} 位成员的公开邮箱(这一步会慢一些) ...")
            for idx, login in enumerate(sorted(members), start=1):
                code, payload = gh.request("GET", f"/users/{urllib.parse.quote(login)}")
                if code < 400 and isinstance(payload, dict) and payload.get("email"):
                    member_emails.add(str(payload["email"]).strip().lower())
                if idx % 50 == 0:
                    log(f"    ... {idx}/{len(members)}")
            log(f"  解析到 {len(member_emails)} 个公开邮箱。")

    # ---- 续跑: 跳过历史已成功处理过的 ----
    done_keys: set[str] = set()
    if not args.no_resume:
        done_keys = load_done_keys(out_dir)
        if done_keys:
            log(f"\n续跑: 读到 {len(done_keys)} 条历史成功记录, 会自动跳过。"
                f"(要全部重跑请加 --no-resume)")

    # ---- 生成计划 ----
    todo: list[Record] = []
    for rec in valid:
        if rec.login and rec.login.lower() in members:
            excluded.append(Outcome(rec, "already_member", detail="已经是组织成员"))
        elif rec.email and rec.email.lower() in member_emails:
            excluded.append(Outcome(rec, "already_member", detail="该邮箱出现在某位成员的公开资料里"))
        elif rec.login and rec.login.lower() in pending_logins:
            excluded.append(Outcome(rec, "already_invited", detail="已有待处理邀请"))
        elif rec.email and rec.email.lower() in pending_emails:
            excluded.append(Outcome(rec, "already_invited", detail="该邮箱已有待处理邀请"))
        elif done_keys and (rec.email.lower() in done_keys or rec.login.lower() in done_keys):
            excluded.append(Outcome(rec, "skipped_resume", detail="上次已成功处理过"))
        else:
            todo.append(rec)

    if args.limit and len(todo) > args.limit:
        log(f"  --limit {args.limit}: 本次只处理前 {args.limit} 条, 其余留到下次。")
        todo = todo[: args.limit]

    log(f"\n本次将邀请 {len(todo)} 人, 跳过 {len(excluded)} 条。")
    for rec in todo:
        who = rec.login or rec.email
        log(f"  → 第 {rec.row} 行  {who}  {rec.email if rec.login else ''}")

    if not args.execute:
        plan_path = out_dir / f"plan-{stamp}.csv"
        with plan_path.open("w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.writer(handle)
            writer.writerow(["row", "email", "username", "name", "action", "detail"])
            for rec in todo:
                writer.writerow([rec.row, rec.email, rec.login, rec.name, "will_invite", rec.note])
            for out in excluded:
                writer.writerow([out.record.row, out.record.email, out.record.login,
                                 out.record.name, out.status, out.detail])
        log(f"\n[DRY-RUN] 没有发送任何邀请。")
        log(f"计划已写入: {plan_path}")
        log("确认无误后, 加 --execute 真正执行(建议先加 --limit 5 小批量试跑)。")
        return 0

    # ---- 真正执行 ----
    team_ids = [int(x) for x in re.split(r"[,\s]+", args.team_ids) if x.strip().isdigit()] if args.team_ids else []
    results_path = out_dir / f"results-{stamp}.csv"
    jsonl_path = out_dir / f"results-{stamp}.jsonl"

    log(f"\n开始发送邀请 (间隔 {args.delay}s) ...\n")
    invite_cache: dict[str, int | None] = {}
    counts: dict[str, int] = {}

    csv_handle = results_path.open("w", newline="", encoding="utf-8-sig")
    csv_writer = csv.writer(csv_handle)
    csv_writer.writerow(["row", "name", "email", "username", "invitee_id",
                         "status", "http_status", "detail", "time"])
    jsonl_handle = jsonl_path.open("w", encoding="utf-8")

    try:
        for index, rec in enumerate(todo, start=1):
            status_code = 0
            detail = ""
            invitee_id: int | None = None
            status = "error"
            body: dict[str, Any] = {"role": args.role}
            if team_ids:
                body["team_ids"] = team_ids

            try:
                if args.prefer_username and rec.login:
                    if rec.login.lower() not in invite_cache:
                        code, payload = gh.request("GET", f"/users/{urllib.parse.quote(rec.login)}")
                        invite_cache[rec.login.lower()] = (
                            int(payload["id"]) if code < 400 and isinstance(payload, dict) and payload.get("id") else None
                        )
                    invitee_id = invite_cache[rec.login.lower()]
                    if invitee_id:
                        body["invitee_id"] = invitee_id
                    elif rec.email:
                        body["email"] = rec.email

                if "invitee_id" not in body:
                    if not rec.email:
                        status = "failed"
                        detail = "无法解析该 GitHub 用户名, 且这一行没有邮箱"
                        code = 0
                        payload = {}
                    else:
                        body["email"] = rec.email
                        code, payload = gh.request("POST", f"/orgs/{args.org}/invitations", body)
                else:
                    code, payload = gh.request("POST", f"/orgs/{args.org}/invitations", body)

                status_code = code
                if 200 <= code < 300:
                    status = "invited"
                    if isinstance(payload, dict) and payload.get("id"):
                        invitee_id = invitee_id or payload.get("id")
                    detail = "邀请已发出"
                elif code == 404:
                    status = "failed"
                    detail = "404: 组织不存在, 或 token 无权访问该组织"
                elif code == 422:
                    raw = _describe(code, payload)
                    lowered = raw.lower()
                    # GitHub 对「已经是成员」会明确回:
                    #   A user with this email address is already a part of this organization
                    # 这是最可靠的判断依据, 比拿成员列表比对邮箱靠谱得多
                    if "already a part of this organization" in lowered \
                            or "already a member" in lowered \
                            or "is already" in lowered:
                        status = "already_member"
                        detail = f"已经是组织成员 ({raw})"
                    elif "already" in lowered or "member" in lowered:
                        status = "already_member"
                        detail = raw
                    else:
                        status = "failed"
                        detail = raw
                else:
                    status = "failed"
                    detail = _describe(code, payload)
            except ApiError as exc:
                status = "error"
                status_code = exc.status
                detail = exc.message
            except Exception as exc:  # noqa: BLE001 - 单条失败不应中断整批
                status = "error"
                detail = f"{type(exc).__name__}: {exc}"

            counts[status] = counts.get(status, 0) + 1
            mark = {"invited": "✓", "already_member": "=", "already_invited": "="}.get(status, "✗")
            log(f"[{index}/{len(todo)}] {mark} 第 {rec.row} 行 {rec.login or rec.email} -> {status}"
                + (f"  ({detail})" if detail and status != "invited" else ""))

            csv_writer.writerow([rec.row, rec.name, rec.email, rec.login, invitee_id or "",
                                 status, status_code, detail, now_iso()])
            csv_handle.flush()
            jsonl_handle.write(json.dumps({
                "row": rec.row, "name": rec.name, "email": rec.email, "login": rec.login,
                "invitee_id": invitee_id, "status": status, "http_status": status_code,
                "detail": detail, "time": now_iso(),
            }, ensure_ascii=False) + "\n")
            jsonl_handle.flush()

            if index < len(todo) and args.delay > 0:
                time.sleep(args.delay)
    finally:
        csv_handle.close()
        jsonl_handle.close()

    log("\n== 汇总 ==")
    for status, count in sorted(counts.items(), key=lambda kv: -kv[1]):
        log(f"  {status}: {count}")
    log(f"\n跳过 {len(excluded)} 条(已存在/重复/无效)。")
    log(f"结果 CSV : {results_path}")
    log(f"结果 JSONL: {jsonl_path}")
    log("\n建议稍等几分钟, 再跑一次对账, 确认谁真收到了、谁本来就已经是成员:")
    log(f"  python {Path(__file__).name} --org {args.org} --input \"{table_path.name}\" --report")
    return 0


def _index_of(header: list[str], name: str) -> int | None:
    target = norm_key(name)
    for idx, col in enumerate(header):
        if norm_key(col) == target:
            return idx
    die(f"表头里找不到列: {name}\n现有表头: {' | '.join(header)}")


def _name_of(header: list[str], idx: int | None) -> str:
    if idx is None or idx >= len(header):
        return "(未使用)"
    return header[idx]


if __name__ == "__main__":
    raise SystemExit(main())
